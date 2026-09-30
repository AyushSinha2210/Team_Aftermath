"""Generate the brief three-page work report using bundled python-docx."""
import argparse
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]


def make_report(commit_count):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(.75)
    section.left_margin = section.right_margin = Inches(.85)
    styles = document.styles
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
        styles[name].font.name = 'Calibri'
        styles[name].font.color.rgb = RGBColor(0, 0, 0)
    styles['Normal'].font.size = Pt(11)
    styles['Normal'].paragraph_format.line_spacing = 1.12
    styles['Normal'].paragraph_format.space_after = Pt(8)
    styles['Title'].font.size = Pt(28)
    styles['Title'].paragraph_format.space_after = Pt(7)
    styles['Subtitle'].font.size = Pt(13)
    styles['Heading 1'].font.size = Pt(16)
    styles['Heading 1'].paragraph_format.space_before = Pt(14)
    styles['Heading 1'].paragraph_format.space_after = Pt(8)
    styles['Heading 2'].font.size = Pt(12)
    document.core_properties.author = 'Abhilash'
    document.core_properties.title = 'Work done'
    document.core_properties.subject = 'Ariadne frontend implementation and validation'

    def p(text, style=None):
        document.add_paragraph(text, style)

    def h(text):
        document.add_heading(text, level=1)

    document.add_heading('Work done', level=0)
    p('Ariadne frontend and API integration', 'Subtitle')
    p('Abhilash  |  30 September 2026')
    p('The Ariadne frontend is implemented and locally validated. It provides a cinematic introduction, transparent retrieval evidence and a working CPU search interface. All new work is saved in the separate Abhilash Prism checkout. GitHub delivery is awaiting authentication as abhi-s99.')
    h('Frontend implementation')
    p('Built the page with React, Vite, Tailwind CSS and Framer Motion. The design uses the specified charcoal background, gold and teal accents, hairline borders and centralized color tokens. Fraunces, Manrope and JetBrains Mono are bundled locally so the interface does not depend on a Google Fonts connection.')
    p('Implemented a generic CSS 3D laptop opening sequence with the shared Hero component inside its screen. The introduction opens automatically, cross-fades into the full page and provides a keyboard-accessible skip control. Reduced-motion preferences and low-power devices skip it; a presentation flag can disable it immediately.')
    p('Added the hero and problem statement, scroll-linked retrieval diagram, once-only section reveals, benchmark count-up, live query interface, comparison table, illustrative function-version diff and minimal PRISM footer. The design-system route exposes the token palette and interaction states for review.')
    h('Repository organization')
    p('The React application lives in frontend. The new HTTP bridge lives in ariadne/frontend_api. Documentation, reproducible scripts, validation notes and the work report are included. Dependencies, build output, caches, screenshots and delivery logs are kept out of Git.')
    p('Working folder: C:/Users/admin/Desktop/PRISM/Abhilash Prism. Branch: Abhilash. Remote: https://github.com/abhi-s99/Team_Aftermath.git.')

    document.add_page_break()
    h('Real retrieval integration')
    p('The repository provided a Streamlit application rather than a browser JSON endpoint. Added a loopback HTTP bridge with a documented, versioned response contract. It uses the existing incremental index and shared fine-tuned CPU encoder, with optional BM25 and reciprocal-rank-fusion mode through the existing HybridPipeline.')
    p('The demonstration searches 15 functions from the checked-in JavaScript voice-assistant sample. Tree-sitter supplies the actual function boundaries. Results show source code, relative file path, start and end lines, and raw retrieval scores. The mini pipeline and attribution tags reflect the mode that actually ran; reranking is not enabled in this demo.')
    h('Input and failure handling')
    p('Implemented empty-input validation, keyboard submission, loading status, duplicate-request guards, cancellation on unmount, an explicit eight-second timeout, calm empty results and inline errors with retry. Source code is rendered as escaped text, and malformed API metadata is rejected before display.')
    p('Captured real API responses for the three supplied example questions. Only those exact questions may use a recording when the backend or proxy is unreachable. The page visibly labels these responses as recorded. Arbitrary offline queries never borrow another example, and timeouts or structured backend errors remain explicit failures.')
    h('Evidence presented accurately')
    p('The official MTEB test result is separate from the smaller validation experiment. The frontend did not rerun either benchmark. Missing latency measurements are shown as not recorded, and the illustrative version edit is labeled accordingly.')
    table = document.add_table(rows=1, cols=4)
    table.autofit = False
    widths = [2.2, 1.6, 1.0, 1.0]
    for cell, width in zip(table.rows[0].cells, widths):
        cell.width = Inches(width)
    for cell, label in zip(table.rows[0].cells, ['Evaluation', 'Scope', 'NDCG at 10', 'MRR at 10']):
        cell.text = label
    for row in [
        ['Official MTEB', 'Test split', '0.159', '0.137'],
        ['Dense fine tuned', '50 validation queries', '0.739', '0.707'],
        ['Dense plus fusion', 'Same validation run', '0.721', '0.682'],
        ['Fusion plus rerank', 'Same validation run', '0.596', '0.531'],
    ]:
        for cell, value in zip(table.add_row().cells, row):
            cell.text = value
    for i, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for edge in ['top', 'left', 'bottom', 'right']:
                element = OxmlElement(f'w:{edge}')
                for key, value in {'val': 'single', 'sz': '5', 'color': 'D9D9D9'}.items():
                    element.set(qn(f'w:{key}'), value)
                borders.append(element)
            tcpr.append(borders)
            shading = OxmlElement('w:shd')
            shading.set(qn('w:fill'), '263445' if i == 0 else ('F2F5F8' if i % 2 == 0 else 'FFFFFF'))
            tcpr.append(shading)
            margins = OxmlElement('w:tcMar')
            for side in ['top', 'left', 'bottom', 'right']:
                element = OxmlElement(f'w:{side}')
                element.set(qn('w:w'), '100')
                element.set(qn('w:type'), 'dxa')
                margins.append(element)
            tcpr.append(margins)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    if i == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)

    document.add_page_break()
    h('Validation completed')
    p('Python API and retrieval-adapter tests: 24 passed. Frontend protocol and search-state tests: 21 passed. Edge browser tests with the real CPU backend enabled: 14 passed. The accessibility scan reported no WCAG A or AA violations in the tested reduced-motion page.')
    p('The browser suite exercised Enter submission, rapid duplicate clicks, loading, empty results, malformed responses, backend and proxy failures, exact-query recordings, arbitrary offline input, timeout and retry, intro handoff, reduced motion, keyboard focus and locally bundled fonts. Layout checks passed at widths of 1280, 1366 and 390 pixels, and browser screenshots were visually inspected.')
    p('The production build passed. Its main JavaScript bundle is about 122 kB gzipped and CSS about 24 kB gzipped. The dependency audit reported zero vulnerabilities. A source audit passed for token usage, absence of UI emoji and shadows, and transform/opacity motion. Tree-sitter emitted only upstream deprecation warnings during Python checks.')
    h('Running the demonstration')
    p('Use Python 3.11 and Node 24. Install the API dependencies from ariadne/frontend_api/requirements.txt, then run python -m ariadne.frontend_api.server from the repository root. In a second Bash terminal, enter frontend, run npm ci and npm run dev, then open http://127.0.0.1:5173. Full commands are in frontend/README.md.')
    p('Use /design-system to inspect tokens. Use /?intro=off to disable the opening immediately. A steady 60 frames per second has not been measured on the final presentation hardware; check that machine before presenting. The configured GitHub Actions workflow has not run remotely while delivery is pending.')
    h('Git delivery status')
    p(f'{commit_count} incremental commits are prepared on Abhilash, authored as abhi-s99 with abhilashsingh2005@gmail.com. Each completed unit goes through the Bash delivery script. The current GitHub login is atharvasheersh and has no push permission on the destination, so the commits are retained in an ordered local queue.')
    p('After signing in as abhi-s99 through GitHub CLI, run bash scripts/deliver.sh --flush from the repository root. It pushes every queued commit in order, records each success, and preserves remaining work if a push fails. The submission release tag is unchanged.')
    document.save(ROOT / 'work done.docx')
    print(f"Created {ROOT / 'work done.docx'}")


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--commit-count', type=int, required=True)
    make_report(parser.parse_args().commit_count)
